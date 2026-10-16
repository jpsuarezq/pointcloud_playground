"""Build a consolidated Excel workbook from the per-tree table:
biomass (AGB), carbon and CO2-equivalent added per tree.

Equation and constants are defined in the CONFIG block below; they are set from
the primary-source research in docs/research/biomass-allometry/README.md.
"""
import csv, os, numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

O = "/tmp/opencode/dbh/out"
SRC = f"{O}/tabla_arboles.csv"
OUT = "/home/usuario/workspaces/NubePuntos/prototype/stem-isolation-dbh/results/biomasa_carbono.xlsx"

# ---------------------------------------------------------------- CONFIG
RHO = 0.60          # wood specific gravity, g/cm3 (species unknown -> default)
CARBON_FRACTION = 0.47   # IPCC default
CO2_PER_C = 44.0 / 12.0  # molecular ratio


def agb_chave2014_dh(d_cm, h_m, rho=RHO):
    """Chave et al. (2014) pantropical, DBH + height. Returns kg."""
    return 0.0673 * (rho * d_cm ** 2 * h_m) ** 0.976


def agb_dbh_only(d_cm, rho=RHO):
    """Chave et al. (2005) moist forest, DBH only (fallback). Returns kg."""
    return rho * np.exp(-1.499 + 2.148 * np.log(d_cm)
                        + 0.207 * (np.log(d_cm)) ** 2
                        - 0.0281 * (np.log(d_cm)) ** 3)


def f(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def main():
    rows = list(csv.DictReader(open(SRC)))
    for r in rows:
        d = f(r["DBH_cm"]); h = f(r["altura_m"])
        if r["estado"] == "measured" and d and h and d > 0 and h > 0:
            agb = float(agb_chave2014_dh(d, h))
            agb_d = float(agb_dbh_only(d))
        else:
            agb = agb_d = None
        r["biomasa_aerea_kg"] = round(agb, 1) if agb is not None else ""
        r["biomasa_dbh_only_kg"] = round(agb_d, 1) if agb_d is not None else ""
        r["carbono_kg"] = round(agb * CARBON_FRACTION, 1) if agb is not None else ""
        r["CO2eq_kg"] = round(agb * CARBON_FRACTION * CO2_PER_C, 1) if agb is not None else ""

    wb = openpyxl.Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="2F5496")

    # ---- sheet: Arboles
    ws = wb.active; ws.title = "Arboles"
    cols = ["id_arbol", "latitud", "longitud", "DBH_cm", "altura_m", "confianza",
            "estado", "arco", "seed", "inclinacion_deg", "contrafuerte", "n_pts",
            "pts_banda", "biomasa_aerea_kg", "biomasa_dbh_only_kg", "carbono_kg",
            "CO2eq_kg", "motivo"]
    ws.append(cols)
    for c in ws[1]:
        c.font = hdr_font; c.fill = hdr_fill
    for r in rows:
        ws.append([r.get(k, "") for k in cols])
    ws.freeze_panes = "A2"
    for i, w in enumerate([9, 11, 12, 8, 9, 10, 9, 7, 10, 13, 11, 8, 9, 16, 18, 11, 11, 16], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # ---- sheet: Resumen
    ws2 = wb.create_sheet("Resumen")
    meas = [r for r in rows if r["estado"] == "measured"]
    def tot(rs, k):
        return round(sum(float(r[k]) for r in rs if r[k] not in ("", None)), 1)
    ws2.append(["Resumen (nube: segmentado_crop.laz, 72.16M pts)"])
    ws2["A1"].font = Font(bold=True, size=13)
    ws2.append([])
    ws2.append(["Arboles de dosel detectados", len(rows)])
    ws2.append(["Arboles medidos", len(meas)])
    ws2.append(["Arboles rechazados", len(rows) - len(meas)])
    ws2.append([])
    ws2.append(["Parametros de la ecuacion"])
    ws2.append(["Ecuacion", "Chave et al. (2014) pantropical, DBH+altura"])
    ws2.append(["Densidad de madera rho (g/cm3)", RHO])
    ws2.append(["Fraccion de carbono", CARBON_FRACTION])
    ws2.append(["CO2/C", round(CO2_PER_C, 4)])
    ws2.append([])
    ws2.append(["Biomasa aerea total (kg)", tot(meas, "biomasa_aerea_kg")])
    ws2.append(["Carbono total (kg)", tot(meas, "carbono_kg")])
    ws2.append(["CO2eq total (kg)", tot(meas, "CO2eq_kg")])
    ws2.append(["CO2eq total (t)", round(tot(meas, "CO2eq_kg") / 1000, 2)])
    ws2.append([])
    ws2.append(["Por confianza", "n", "biomasa_kg", "carbono_kg", "CO2eq_kg", "CO2eq_t"])
    for c in ws2[ws2.max_row]:
        c.font = Font(bold=True)
    for conf in ("high", "medium", "low"):
        rs = [r for r in meas if r["confianza"] == conf]
        ws2.append([conf, len(rs), tot(rs, "biomasa_aerea_kg"), tot(rs, "carbono_kg"),
                    tot(rs, "CO2eq_kg"), round(tot(rs, "CO2eq_kg") / 1000, 3)])
    for i, w in enumerate([34, 12, 14, 14, 14, 12], 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # ---- sheet: Metodologia
    ws3 = wb.create_sheet("Metodologia")
    txt = [
        "Biomasa y carbono equivalente",
        "",
        "Ecuacion (Chave et al. 2014, pantropical, DBH + altura):",
        "  AGB [kg] = 0.0673 * (rho * DBH[cm]^2 * H[m])^0.976",
        "  con rho = densidad de madera especifica en g/cm3.",
        "",
        f"rho (especie desconocida) = {RHO} g/cm3  [fuente: docs/research/biomass-allometry/README.md]",
        "Fallback (solo DBH, Chave et al. 2005, bosque humedo):",
        "  AGB [kg] = rho * exp(-1.499 + 2.148 ln D + 0.207 (ln D)^2 - 0.0281 (ln D)^3)",
        "",
        f"Carbono = AGB * {CARBON_FRACTION}   (fraccion de carbono por defecto IPCC)",
        f"CO2eq = Carbono * 44/12 = Carbono * {round(CO2_PER_C,4)}",
        "",
        "Advertencias:",
        "- Sin verdad de campo ni especie: rho es un valor por defecto; AGB escala ~lineal con rho.",
        "- Los errores de la ecuacion (tipo destructivo) y los del DBH/altura fotogrametricos se propagan.",
        "- Solo se calcula biomasa para arboles 'measured'; los 'refused' quedan vacios.",
        "- Filas 'low' son de baja confianza: reportar estratificado, no como un unico numero.",
    ]
    for line in txt:
        ws3.append([line])
    ws3["A1"].font = Font(bold=True, size=13)
    ws3.column_dimensions["A"].width = 110

    wb.save(OUT)
    print("wrote", OUT)
    print("measured", len(meas), "AGB total kg", tot(meas, "biomasa_aerea_kg"),
          "CO2eq t", round(tot(meas, "CO2eq_kg") / 1000, 2))


if __name__ == "__main__":
    main()
